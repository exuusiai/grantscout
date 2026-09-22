"""DOCX export for proposal documents (python-docx based)."""

from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Pt

from grantscout.proposal.renderer import _references, _renumber
from grantscout.proposal.schemas import ProposalDocument


def _add_markdown_lines(docx, text: str) -> None:
    """Write light markdown (paragraphs, '- ' bullets, '> ' notes) into docx."""
    for block in text.split("\n"):
        line = block.strip()
        if not line:
            continue
        if line.startswith("- "):
            docx.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("> "):
            paragraph = docx.add_paragraph()
            paragraph.add_run(line[2:]).italic = True
        else:
            docx.add_paragraph(line)


def export_proposal_docx(document: ProposalDocument, output_path: Path) -> Path:
    """Render the proposal as a .docx file and return the written path."""
    docx = DocxDocument()
    docx.add_heading(document.title, level=0)

    docx.add_heading("选题简报", level=1)
    brief = document.idea_brief
    docx.add_paragraph(f"选题:{brief.topic}")
    if brief.goals:
        docx.add_paragraph("研究目标:" + ";".join(brief.goals))
    if brief.innovation_points:
        docx.add_paragraph("创新点:" + ";".join(brief.innovation_points))
    if brief.duration_months:
        docx.add_paragraph(f"研究周期:{brief.duration_months} 个月")
    for question in brief.open_questions:
        docx.add_paragraph(f"待澄清:{question}", style="List Bullet")

    if document.dossier is not None:
        docx.add_heading("研究空白与科学问题", level=1)
        for gap in document.dossier.gaps:
            marker = "" if gap.evidence_ids else "(无证据,待人工补充)"
            docx.add_paragraph(f"空白:{gap.statement}{marker}", style="List Bullet")
        for science_question in document.dossier.science_questions:
            docx.add_paragraph(f"科学问题:{science_question.question}", style="List Bullet")

    docx.add_heading("正文草稿", level=1)
    references, mappings = _references(document)
    spec_titles = {}
    document_outline = document.outline.sections if document.outline else []
    outline_by_key = {section.key: section for section in document_outline}
    for section in document_outline:
        spec_titles[section.key] = section.title
    for section_key, versions in document.sections.items():
        draft = versions[-1] if versions else None
        if draft is None:
            continue
        title = spec_titles.get(section_key, section_key)
        docx.add_heading(f"{title}", level=2)
        _add_markdown_lines(docx, _renumber(draft.content_md, mappings.get(section_key, {})))
        if draft.indicators:
            table = docx.add_table(rows=1, cols=6)
            table.style = "Table Grid"
            header = table.rows[0].cells
            for cell, text in zip(
                header,
                ("指标名称", "目标值", "测试条件", "测试方法", "验收材料", "来源依据"),
            ):
                cell.text = text
            for indicator in draft.indicators:
                row = table.add_row().cells
                values = (
                    indicator.name,
                    indicator.target_value,
                    indicator.test_conditions,
                    indicator.test_method,
                    indicator.acceptance_materials,
                    indicator.source_basis,
                )
                for cell, value in zip(row, values):
                    cell.text = value or "-"
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
        if outline_by_key.get(section_key) is not None and not document.outline.approved:
            docx.add_paragraph("(大纲未经人工确认,内容为草稿)", style="Intense Quote")

    if references:
        docx.add_heading("参考文献", level=1)
        for number, (evidence_id, paper_title) in enumerate(references, start=1):
            docx.add_paragraph(f"[{number}] {paper_title}(证据 {evidence_id})")

    if document.warnings:
        docx.add_heading("全局提示", level=1)
        for warning in document.warnings:
            docx.add_paragraph(warning, style="List Bullet")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    docx.save(str(output_path))
    return output_path
