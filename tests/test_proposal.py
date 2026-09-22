import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from typer.testing import CliRunner

from grantscout.config import Settings
from grantscout.proposal import ProposalPipeline, ProposalRequest
from grantscout.proposal.schemas import IndicatorItem
from grantscout.proposal.renderer import render_proposal_markdown
from grantscout.proposal.templates import get_template, list_templates
from grantscout.retrieval.parser import parse_document
from grantscout.retrieval.store import CorpusStore

PAPER_ONE = (
    "Introduction\n\n"
    "Network emulation platforms such as EMANE and ns-3 are widely used for satellite "
    "network research.\n\n"
    "Method\n\n"
    "We propose StarryNet, a container-based network emulator for LEO satellite "
    "constellations.\n\n"
    "Results\n\n"
    "StarryNet improves emulation scalability by 40 percent on the benchmark.\n\n"
    "Limitations\n\n"
    "However, the emulator fails to model inter-satellite laser links accurately.\n"
)

PAPER_TWO = (
    "Introduction\n\n"
    "Digital twin systems for satellite networks require high-fidelity ground station "
    "emulation.\n\n"
    "Method\n\n"
    "We introduce a hybrid emulation framework that couples real gateways with the ns-3 "
    "discrete event simulator.\n\n"
    "Limitations\n\n"
    "However, the framework lacks support for dynamic topology reconfiguration during "
    "emulation runs.\n"
)

TOPIC = "StarryNet network emulation platform for satellite constellations"


class FakeProposalModel:
    """Doubles as the gap/outline/indicator JSON model and the drafting model."""

    def __init__(self, gaps=None, outline=None, content="", indicators=None):
        self.gaps = gaps or {"gaps": []}
        self.outline = outline or {"sections": []}
        self.content = content
        self.indicators = indicators or {"content_md": content, "indicators": []}

    def chat_json(self, *, messages, max_tokens, temperature):
        joined = messages[0]["content"] + "\n" + messages[-1]["content"]
        if '"indicators"' in joined:
            return self.indicators
        if '"sections"' in joined:
            return self.outline
        return self.gaps

    def chat(self, *, messages, max_tokens, temperature):
        return SimpleNamespace(content=self.content, model="fake", usage={})


def _settings(tmp_path: Path) -> Settings:
    return Settings(
        data_dir=tmp_path / "data",
        runs_dir=tmp_path / "runs",
        use_model_reasoning=False,
        max_steps=40,
        max_tool_calls=60,
        max_papers=5,
    )


def _build_corpus(tmp_path: Path) -> Path:
    corpus_path = tmp_path / "corpus.sqlite"
    with CorpusStore(corpus_path) as store:
        for paper_id, text in (("starrynet", PAPER_ONE), ("hybrid", PAPER_TWO)):
            source = tmp_path / f"{paper_id}.txt"
            source.write_text(text, encoding="utf-8")
            store.upsert(parse_document(source, paper_id=paper_id))
    return corpus_path


def test_template_registry_lists_builtin_and_rejects_unknown() -> None:
    ids = {template.id for template in list_templates()}
    assert "project-plan-v1" in ids
    template = get_template("project-plan-v1")
    assert [spec.key for spec in template.sections] == [
        "requirements",
        "state-of-the-art",
        "system-design",
        "technical-approach",
        "metrics",
    ]
    with pytest.raises(KeyError):
        get_template("missing-template")


def test_pipeline_deterministic_run_produces_traceable_draft(tmp_path: Path) -> None:
    corpus_path = _build_corpus(tmp_path)
    settings = _settings(tmp_path)
    with CorpusStore(corpus_path) as store:
        pipeline = ProposalPipeline(store, settings, output_language="zh")
        document = pipeline.run(ProposalRequest(topic=TOPIC, goals=["可扩展仿真"]))

    assert document.status == "completed"
    assert document.dossier is not None
    assert document.dossier.gaps, "limitations in the corpus must become gap statements"
    assert all(gap.evidence_ids for gap in document.dossier.gaps)
    assert document.outline is not None
    assert set(document.sections) == {spec.key for spec in get_template("project-plan-v1").sections}
    citations = [c for versions in document.sections.values() for c in versions[-1].citations]
    assert citations, "retrieved evidence must surface as citations"
    assert document.review is not None
    assert document.review.total_citations == len(citations)
    assert (tmp_path / "runs" / f"{document.id}.proposal.json").exists()
    assert (tmp_path / "runs" / f"{document.id}.proposal.docx").exists()
    report = render_proposal_markdown(document)
    assert "参考文献" in report
    assert "【文献" in report


def test_pipeline_degrades_to_placeholders_without_evidence(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    empty_corpus = tmp_path / "empty.sqlite"
    with CorpusStore(empty_corpus):
        pass
    with CorpusStore(empty_corpus) as store:
        pipeline = ProposalPipeline(store, settings, output_language="zh")
        document = pipeline.run(ProposalRequest(topic="completely unrelated topic xyz"))

    assert document.status == "completed"
    assert document.dossier is not None
    assert document.dossier.gaps[0].confidence == 0.0
    for spec in get_template("project-plan-v1").sections:
        draft = document.latest_draft(spec.key)
        assert draft is not None
        assert "【待补充】" in draft.content_md
    assert document.review is not None
    assert set(document.review.sections_missing_evidence) == {
        spec.key for spec in get_template("project-plan-v1").sections
    }


def test_pipeline_model_path_validates_citations_and_gaps(tmp_path: Path) -> None:
    corpus_path = _build_corpus(tmp_path)
    settings = _settings(tmp_path)
    with CorpusStore(corpus_path) as store:
        evidence_ids = sorted(item.id for _, item in store.list_evidence())
        valid_id = evidence_ids[0]
        fake = FakeProposalModel(
            gaps={
                "gaps": [
                    {
                        "statement": "Laser link modeling remains unsolved.",
                        "evidence_ids": [valid_id, "invented:evidence"],
                    }
                ]
            },
            outline={
                "sections": [
                    {"key": "requirements", "bullet_points": ["b1", "b2"], "evidence_ids": [valid_id]}
                ]
            },
            content=(
                "<think>内部推理,不应出现在正文。</think>"
                "容器化仿真路线已有实证支撑【文献1】。性能提升显著【文献1】。"
                "激光链路建模仍为空白【文献9】。"
            ),
        )
        pipeline = ProposalPipeline(store, settings, output_language="zh")
        pipeline.model_client = fake
        document = pipeline.run(ProposalRequest(topic=TOPIC))

    assert document.dossier is not None
    assert document.dossier.gaps[0].evidence_ids == [valid_id]
    assert document.dossier.gaps[0].confidence == 0.8
    requirements = document.latest_draft("requirements")
    assert requirements is not None and requirements.source == "agent"
    assert "【文献9】" not in requirements.content_md
    assert "内部推理" not in requirements.content_md
    assert [citation.marker for citation in requirements.citations] == ["【文献1】"]
    assert requirements.citations[0].evidence_id in evidence_ids
    assert any("剔除非法引用" in warning for warning in requirements.warnings)
    outline_section = next(
        item for item in (document.outline.sections if document.outline else []) if item.key == "requirements"
    )
    assert outline_section.bullet_points == ["b1", "b2"]
    assert outline_section.evidence_ids == [valid_id]


def test_cli_propose_writes_artifacts(tmp_path: Path, monkeypatch) -> None:
    from grantscout import cli

    corpus_path = _build_corpus(tmp_path)
    settings = _settings(tmp_path)
    monkeypatch.setattr(cli, "get_settings", lambda: settings)
    result = CliRunner().invoke(
        cli.app, ["propose", TOPIC, "--corpus", str(corpus_path)]
    )

    assert result.exit_code == 0, result.output
    payload = json.loads(result.output)
    assert payload["status"] == "completed"
    assert Path(payload["artifacts"]["markdown"]).exists()


def test_propose_reports_missing_corpus(tmp_path: Path, monkeypatch) -> None:
    from grantscout import cli

    settings = _settings(tmp_path)
    monkeypatch.setattr(cli, "get_settings", lambda: settings)
    result = CliRunner().invoke(
        cli.app, ["propose", TOPIC, "--corpus", str(tmp_path / "missing.sqlite")]
    )

    assert result.exit_code == 1
    assert "does not exist" in result.output


def test_real_templates_are_registered_with_indicator_sections() -> None:
    ids = {template.id for template in list_templates()}
    assert {
        "project-plan-v1",
        "universal-plan-v1",
        "lasa-key-project-2023",
        "most-innovation-2015",
        "jiangsu-erc-2023",
        "guangxi-collab-2025",
    } <= ids
    lasa = get_template("lasa-key-project-2023")
    structured = {spec.key for spec in lasa.sections if spec.structured_indicators}
    assert structured == {"technical-indicators"}
    manual = {spec.key for spec in lasa.sections if not spec.generatable}
    assert {"budget-form", "personnel-form", "ethics-form"} <= manual
    assert len(get_template("universal-plan-v1").sections) == 14
    six_fields = {"name", "target_value", "test_conditions", "test_method", "acceptance_materials", "source_basis"}
    assert set(IndicatorItem.model_fields) == six_fields


def test_pipeline_structured_indicators_and_manual_sections(tmp_path: Path) -> None:
    corpus_path = _build_corpus(tmp_path)
    settings = _settings(tmp_path)
    with CorpusStore(corpus_path) as store:
        evidence_ids = sorted(item.id for _, item in store.list_evidence())
        valid_id = evidence_ids[0]
        fake = FakeProposalModel(
            gaps={
                "gaps": [
                    {"statement": "Laser link modeling remains unsolved.", "evidence_ids": [valid_id]}
                ]
            },
            outline={
                "sections": [
                    {
                        "key": "technical-indicators",
                        "bullet_points": ["量化指标"],
                        "evidence_ids": [valid_id],
                    }
                ]
            },
            content="基于资料梳理考核指标如下。",
            indicators={
                "content_md": "考核指标基于资料整理【文献1】。",
                "indicators": [
                    {
                        "name": "证据召回率",
                        "target_value": "≥ 0.8",
                        "test_conditions": "标准测试语料库",
                        "test_method": "Recall@10 自动评测",
                        "acceptance_materials": "评测报告",
                        "source_basis": "【文献1】",
                    },
                    {"name": "", "target_value": "x"},
                ],
            },
        )
        pipeline = ProposalPipeline(
            store, settings, template_id="lasa-key-project-2023", output_language="zh"
        )
        pipeline.model_client = fake
        document = pipeline.run(
            ProposalRequest(topic=TOPIC, template_id="lasa-key-project-2023")
        )

    indicators_draft = document.latest_draft("technical-indicators")
    assert indicators_draft is not None and indicators_draft.source == "agent"
    assert len(indicators_draft.indicators) == 1
    indicator = indicators_draft.indicators[0]
    assert indicator.name == "证据召回率"
    assert indicator.test_method == "Recall@10 自动评测"
    assert "【文献1】" in {citation.marker for citation in indicators_draft.citations}
    budget_form = document.latest_draft("budget-form")
    assert budget_form is not None and "人工填写" in budget_form.content_md
    assert "budget-form" not in (document.review.sections_missing_evidence if document.review else [])
    report = render_proposal_markdown(document)
    assert "| 指标名称 | 目标值 | 测试条件 | 测试方法 | 验收材料 | 来源依据 |" in report
    assert "Recall@10 自动评测" in report


def test_docx_export_is_readable(tmp_path: Path) -> None:
    from docx import Document as DocxDocument

    from grantscout.proposal.exporter import export_proposal_docx

    corpus_path = _build_corpus(tmp_path)
    settings = _settings(tmp_path)
    settings = settings.model_copy(update={"proposal_expansion_passes": 0})
    with CorpusStore(corpus_path) as store:
        pipeline = ProposalPipeline(store, settings, output_language="zh")
        document = pipeline.run(ProposalRequest(topic=TOPIC))

    docx_path = tmp_path / "out" / "proposal.docx"
    export_proposal_docx(document, docx_path)
    docx = DocxDocument(str(docx_path))
    text = "\n".join(paragraph.text for paragraph in docx.paragraphs)
    assert document.title in text
    assert "正文草稿" in text
    assert "参考文献" in text
    headings = [paragraph.text for paragraph in docx.paragraphs if paragraph.style.name.startswith("Heading")]
    assert any("技术路线" in heading or "国内外研究现状" in heading for heading in headings)


def test_proposal_store_roundtrip_and_gates(tmp_path: Path) -> None:
    from grantscout.proposal.store import ProposalStore

    corpus_path = _build_corpus(tmp_path)
    settings = _settings(tmp_path)
    runs_dir = tmp_path / "runs"
    with CorpusStore(corpus_path) as store:
        pipeline = ProposalPipeline(store, settings, output_language="zh")
        document = pipeline.run(ProposalRequest(topic=TOPIC, template_id="lasa-key-project-2023"))

        proposal_store = ProposalStore(runs_dir)
        assert [item["id"] for item in proposal_store.list()] == [document.id]

        document.outline.approved = True
        proposal_store.save(document)
        reloaded = proposal_store.get(document.id)
        assert reloaded.outline.approved is True

        section_key = "goals-tasks-analysis"
        before = len(reloaded.sections[section_key])
        updated = pipeline.redraft_section(reloaded, section_key)
        assert len(updated.sections[section_key]) == before + 1
        assert updated.sections[section_key][-1].version == before + 1
        with pytest.raises(ValueError):
            pipeline.redraft_section(updated, "budget-form")
    with pytest.raises(KeyError):
        ProposalStore(runs_dir).get("missing-id")


def test_api_proposal_gates_and_downloads(tmp_path: Path, monkeypatch) -> None:
    from grantscout.api import app as app_module

    _build_corpus(tmp_path / "data")
    settings = _settings(tmp_path)
    monkeypatch.setattr(app_module, "get_settings", lambda: settings)
    client = TestClient(app_module.app)

    created = client.post("/api/proposals/draft", json={"topic": TOPIC})
    assert created.status_code == 200, created.text
    proposal_id = created.json()["document"]["id"]

    listed = client.get("/api/proposals").json()
    assert proposal_id in {item["id"] for item in listed}

    approved = client.post(
        f"/api/proposals/{proposal_id}/outline/approve", json={"approved": True}
    )
    assert approved.status_code == 200
    assert approved.json()["document"]["outline"]["approved"] is True

    regenerated = client.post(f"/api/proposals/{proposal_id}/sections/requirements/regenerate")
    assert regenerated.status_code == 200, regenerated.text
    versions = regenerated.json()["document"]["sections"]["requirements"]
    assert len(versions) == 2 and versions[-1]["version"] == 2

    manual = client.post(f"/api/proposals/{proposal_id}/sections/overview/regenerate")
    assert manual.status_code == 422  # project-plan-v1 has no 'overview'; unknown section
    unknown = client.get("/api/proposals/does-not-exist")
    assert unknown.status_code == 404

    markdown = client.get(f"/api/proposals/{proposal_id}/report.md")
    assert markdown.status_code == 200 and "正文草稿" in markdown.text
    docx_download = client.get(f"/api/proposals/{proposal_id}/report.docx")
    assert docx_download.status_code == 200
    assert docx_download.content[:2] == b"PK"  # docx is a zip container
    page = client.get("/proposals")
    assert page.status_code == 200 and "本子工作台" in page.text


def test_api_proposal_endpoint_enforces_data_dir(tmp_path: Path, monkeypatch) -> None:
    from grantscout.api import app as app_module

    data_dir = tmp_path / "data"
    corpus_path = _build_corpus(data_dir)
    settings = _settings(tmp_path)
    monkeypatch.setattr(app_module, "get_settings", lambda: settings)
    client = TestClient(app_module.app)

    response = client.post("/api/proposals/draft", json={"topic": TOPIC})
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["document"]["status"] == "completed"
    assert "正文草稿" in payload["report"]

    outside = client.post(
        "/api/proposals/draft",
        json={"topic": TOPIC, "corpus": str(tmp_path / "outside.sqlite")},
    )
    assert outside.status_code == 422

    empty = client.post(
        "/api/proposals/draft",
        json={"topic": TOPIC, "corpus": str(tmp_path / "data" / "empty.sqlite")},
    )
    assert empty.status_code == 404
    assert corpus_path.exists()
